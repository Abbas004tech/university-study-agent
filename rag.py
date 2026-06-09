from pptx import Presentation
from pypdf import PdfReader


def extract_text_from_pptx(file_path):
    prs = Presentation(file_path)

    text = ""

    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text += shape.text + "\n"

    return text


def extract_text_from_pdf(file_path):
    reader = PdfReader(file_path)

    text = ""

    for page in reader.pages:
        page_text = page.extract_text()

        if page_text:
            text += page_text + "\n"

    return text


def chunk_text(text, chunk_size=4000):

    chunks = []

    for i in range(0, len(text), chunk_size):
        chunks.append(text[i:i + chunk_size])

    return chunks
